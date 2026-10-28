"""Recall sweep, long-contract map/reduce, the AI-call budget, client-side inference, score calibration and inputs."""
import io
import json
import re
from types import SimpleNamespace

from app.core.rag.parsing import parse_pdf_pages, parse_text_bytes
from app.workspaces.contract_intelligence import service
from app.workspaces.contract_intelligence.engine import inputs, pipeline, scoring
from app.workspaces.contract_intelligence.engine.schemas import (
    ChecklistSweep, ContractFacts, PartReview, RiskAnalysis, Severity,
)
from app.workspaces.contract_intelligence.engine.segmenter import segment_contract

LEASE = """LEAVE AND LICENCE AGREEMENT
This agreement is made at Surat between Mr. Ramesh Patel ("Licensor") and Surat Textiles Private Limited ("Licensee").
1. TERM
The licence is for thirty-six (36) months from 1 April 2026 for use of the shop premises at Ring Road.
2. LOCK-IN
The Licensee agrees that the entire term of thirty-six (36) months shall be a lock-in period during which it cannot terminate.
3. LICENCE FEE
The monthly licence fee is Rs 1,00,000 plus GST. The licence fee shall be increased by fifteen per cent (15%) every twelve (12) months.
4. SECURITY DEPOSIT
The Licensor shall refund the security deposit within one hundred and eighty (180) days of handover of the premises.
5. TERMINATION
The Licensor may terminate this agreement at any time by giving fifteen (15) days' notice to the Licensee.
6. REPAIRS
The Licensee shall carry out all repairs including structural repairs to the building at its own cost.
7. INSPECTION
The Licensor may inspect the premises at any time with or without notice to the Licensee.
8. GOVERNING LAW
This agreement is governed by the laws of India and the courts at Surat shall have exclusive jurisdiction.
9. STAMP DUTY
This agreement shall be registered and the stamp duty and registration charges shall be borne by the Licensee.
"""
Q_REPAIRS = "The Licensee shall carry out all repairs including structural repairs to the building at its own cost."
Q_INSPECT = "The Licensor may inspect the premises at any time with or without notice to the Licensee."


def _doc(text):
    return parse_text_bytes(text.encode("utf-8"))


def _facts(**over):
    d = dict(title="Leave and Licence", contract_type="lease", summary_points=["Shop licence."],
             parties=[{"name": "Ramesh Patel", "role": "Licensor"}, {"name": "Surat Textiles Private Limited", "role": "Licensee"}],
             likely_client="Surat Textiles Private Limited", termination_notice_days_other=15, stamp_duty_mentioned=True)
    d.update(over)
    return json.dumps(d)


def _finding(quote, seg, cat="other", sev="medium", title="Issue"):
    return {"segment_id": seg, "category": cat, "severity": sev, "title": title, "why_it_matters": "x",
            "quote": quote, "recommendation": "y"}


def _schemas(state):
    return [c.schema.__name__ for c in state.calls]


# ------------------------------------------------------------------ recall sweep

def _sweep_handler(call):
    if call.schema is ContractFacts:
        return _facts()
    if call.schema is RiskAnalysis:
        return json.dumps({"findings": [_finding(Q_REPAIRS, "S7", title="Licensee pays even structural repairs")],
                           "missing_protections": []})
    if call.schema is ChecklistSweep:
        return json.dumps({"items": [
            {"topic": "eviction, re-entry and overstay charges", "status": "present_hurts", "category": "other",
             "segment_id": "S8", "quote": Q_INSPECT, "severity": "medium", "title": "Landlord can walk in without notice",
             "why_it_matters": "No privacy.", "recommendation": "Ask for 24 hours' notice."},
            {"topic": "renewal terms", "status": "present_hurts", "category": "other", "segment_id": "S2",
             "quote": "The Licensee shall have no right of renewal under any circumstances whatsoever.",
             "severity": "high", "title": "No renewal right", "why_it_matters": "x", "recommendation": "y"},
            {"topic": "repairs and maintenance", "status": "present_hurts", "category": "other", "segment_id": "S7",
             "quote": Q_REPAIRS, "severity": "medium", "title": "All repairs on you", "why_it_matters": "x", "recommendation": "y"},
            {"topic": "subletting and change of control", "status": "absent", "category": "assignment_subcontract",
             "title": "No right to sublet or share the premises", "why_it_matters": "You cannot share space with a group company."},
            {"topic": "stamp duty and registration (over 11 months)", "status": "present_fair", "category": "stamp_execution",
             "segment_id": "S10", "quote": "This agreement shall be registered and the stamp duty"},
        ]})
    raise AssertionError(call.schema)


def test_sweep_runs_once_when_few_findings_and_adds_only_new_verified_issues(fake_llm):
    state = fake_llm(_sweep_handler)
    res = pipeline.analyze_parsed(_doc(LEASE), org_name="Some Other Org Pvt Ltd")
    assert _schemas(state) == ["ContractFacts", "RiskAnalysis", "ChecklistSweep"]          # budget: 3 calls
    assert res.review["ai_calls"] == 3 and res.review["sweep_run"] and res.review["sweep_added"] == 1
    titles = {f.title: f for f in res.findings}
    assert titles["Landlord can walk in without notice"].verified and titles["Landlord can walk in without notice"].counts_toward_score
    assert "No renewal right" not in titles                       # invented quote: not added
    assert "All repairs on you" not in titles                     # same clause, same issue as the first pass
    assert titles["No right to sublet or share the premises"].source == "missing"
    assert "1. lock-in period and exit penalty" in state.calls[2].prompt    # topics are walked one by one


def test_client_is_inferred_when_the_organisation_is_not_a_party(fake_llm):
    state = fake_llm(_sweep_handler)
    res = pipeline.analyze_parsed(_doc(LEASE), org_name="Some Other Org Pvt Ltd")
    assert [p.role for p in res.facts.parties if p.is_our_client] == ["Licensee"]
    assert "Surat Textiles Private Limited (Licensee) - assumed" in state.calls[1].prompt


def test_client_falls_back_to_the_usual_sme_role_for_the_type(fake_llm):
    def handler(call):
        if call.schema is ContractFacts:
            return _facts(likely_client=None)
        return _sweep_handler(call)
    fake_llm(handler)
    res = pipeline.analyze_parsed(_doc(LEASE))
    assert [p.role for p in res.facts.parties if p.is_our_client] == ["Licensee"]


def test_lease_with_lock_in_escalation_and_one_sided_termination_scores_high(fake_llm):
    fake_llm(_sweep_handler)
    res = pipeline.analyze_parsed(_doc(LEASE), org_name="Some Other Org Pvt Ltd")
    rule_ids = {f.rule_id for f in res.findings if f.rule_id}
    assert {"IN-LOCKIN-24M", "IN-ESCALATION-10", "IN-TERM-ONESIDED", "IN-DEPOSIT-REFUND"} <= rule_ids
    assert res.score.level == "high"
    b = res.score.breakdown()
    assert sum(c["points"] for c in b["contributions"]) >= res.score.score
    assert all(c["title"] != "No renewal right" for c in b["contributions"])


def test_no_sweep_when_the_risk_pass_already_found_enough(fake_llm):
    quotes = [("The Licensee agrees that the entire term of thirty-six (36) months shall be a lock-in period", "S3"),
              ("The licence fee shall be increased by fifteen per cent (15%) every twelve (12) months.", "S4"),
              ("The Licensor shall refund the security deposit within one hundred and eighty (180) days", "S5"),
              ("The Licensor may terminate this agreement at any time by giving fifteen (15) days' notice", "S6"),
              (Q_REPAIRS, "S7")]

    def handler(call):
        if call.schema is ContractFacts:
            return _facts()
        if call.schema is RiskAnalysis:
            return json.dumps({"findings": [_finding(q, s, title=f"Issue {i}") for i, (q, s) in enumerate(quotes)]})
        raise AssertionError(call.schema)
    state = fake_llm(handler)
    res = pipeline.analyze_parsed(_doc(LEASE))
    assert _schemas(state) == ["ContractFacts", "RiskAnalysis"] and not res.review["sweep_run"]


def test_one_page_purchase_order_works_without_extra_calls(fake_llm):
    po = ("PURCHASE ORDER No. PO/2026/114 dated 05/10/2026\nTo: Shree Ganesh Fabricators, Rajkot\n"
          "Item | Qty | Rate | Amount\nMS brackets 40mm | 500 | Rs 42 | Rs 21,000\n"
          "Delivery within 15 days of this order. Payment within 60 days from invoice. Prices are firm and inclusive of GST.")

    def handler(call):
        if call.schema is ContractFacts:
            return json.dumps(dict(title="Purchase order", contract_type="purchase_order", summary_points=["PO."],
                                   parties=[{"name": "Patel Engineering", "role": "Buyer"},
                                            {"name": "Shree Ganesh Fabricators", "role": "Supplier"}], payment_days=60))
        if call.schema is RiskAnalysis:
            return json.dumps({"findings": []})
        raise AssertionError(call.schema)
    state = fake_llm(handler)
    prepared = pipeline.prepare_document(po.encode(), "po.txt")
    res = pipeline.analyze_parsed(prepared.parsed, org_name="Shree Ganesh Fabricators")
    assert len(state.calls) == 2 and res.score is not None
    assert [p.name for p in res.facts.parties if p.is_our_client] == ["Shree Ganesh Fabricators"]


# ------------------------------------------------------------------ long contracts: map / reduce

def _cost(s):
    return len(s.text) + len(s.heading) + 40


def _two_part_budget(text):
    segs = segment_contract(_doc(text))
    total, biggest = sum(_cost(s) for s in segs), max(_cost(s) for s in segs)
    return (total + biggest) // 2 + 10


def test_long_contract_is_reviewed_end_to_end_with_one_extra_call(fake_llm, monkeypatch):
    monkeypatch.setattr(pipeline, "MAX_DOC_CHARS", _two_part_budget(LEASE))

    def handler(call):
        if call.schema is ContractFacts:
            return _facts(stamp_duty_mentioned=None)
        if call.schema is RiskAnalysis:
            return json.dumps({"findings": [], "missing_protections": [
                {"category": "dispute_resolution", "title": "No dispute clause", "why_it_matters": "x"}]})
        if call.schema is PartReview:
            return json.dumps({"findings": [_finding(Q_INSPECT, "S8", title="Landlord can walk in without notice")],
                               "facts_update": {"governing_law": "Laws of India", "jurisdiction": "Courts at Surat",
                                                "stamp_duty_mentioned": True}})
        raise AssertionError(call.schema)
    state = fake_llm(handler)
    res = pipeline.analyze_parsed(_doc(LEASE))
    assert _schemas(state) == ["ContractFacts", "RiskAnalysis", "PartReview"]           # no sweep: budget used
    part_prompt = state.calls[2].prompt
    assert "part 2 of 2" in part_prompt and "[[S1 |" not in part_prompt and "INSPECTION" in part_prompt
    assert "part 1 of 2" in state.calls[1].prompt and "INSPECTION" not in state.calls[1].prompt
    assert res.coverage.parts == 2 and not res.coverage.truncated
    assert res.facts.governing_law == "Laws of India"                   # fact only the second part states
    assert any(f.title == "Landlord can walk in without notice" and f.verified for f in res.findings)
    assert not any(f.source == "missing" for f in res.findings)        # AI absence claims skipped for split reviews
    assert any("reviewed in 2 parts" in w for w in res.warnings)


def test_text_beyond_two_parts_is_reported_and_absence_rules_skipped(fake_llm, monkeypatch):
    monkeypatch.setattr(pipeline, "MAX_DOC_CHARS", 250)

    def handler(call):
        if call.schema is ContractFacts:
            return _facts(stamp_duty_mentioned=False)
        if call.schema is RiskAnalysis:
            return json.dumps({"findings": []})
        if call.schema is PartReview:
            return json.dumps({"findings": []})
        raise AssertionError(call.schema)
    fake_llm(handler)
    res = pipeline.analyze_parsed(_doc(LEASE))
    assert res.coverage.truncated and res.coverage.parts == 2
    assert not any(f.rule_id in pipeline.ABSENCE_RULES for f in res.findings)
    assert any("very long" in w for w in res.warnings)


# ------------------------------------------------------------------ evidence and scoring

def test_unverified_ai_finding_never_counts_even_when_a_rule_agrees(fake_llm):
    def handler(call):
        if call.schema is ContractFacts:
            return _facts()
        if call.schema is RiskAnalysis:
            return json.dumps({"findings": [_finding("The Licensor can end this licence whenever it wishes without any reason.",
                                                     "S6", cat="termination", sev="high", title="Licensor can end it anytime")]})
        if call.schema is ChecklistSweep:
            return json.dumps({"items": []})
        raise AssertionError(call.schema)
    fake_llm(handler)
    res = pipeline.analyze_parsed(_doc(LEASE))
    term = [f for f in res.findings if f.category == "termination"]
    assert len(term) == 1 and term[0].source == "rule" and term[0].rule_id == "IN-TERM-ONESIDED"
    assert term[0].verified and "fifteen (15) days" in term[0].quote
    assert all(f.verified for f in res.findings if f.counts_toward_score and f.source == "ai")


def test_score_breakdown_lists_contributions_and_excludes_unverified():
    res = scoring.compute_score([
        scoring.ScoreInput("lock_in_exclusivity", Severity.high, True, "Lock-in", "rule", "IN-LOCKIN-24M"),
        scoring.ScoreInput("payment_terms", Severity.high, True, "Escalation", "rule", "IN-ESCALATION-10"),
        scoring.ScoreInput("termination", Severity.high, True, "One-sided exit", "rule", "IN-TERM-ONESIDED"),
        scoring.ScoreInput("other", Severity.critical, False, "Unverified claim"),
    ])
    b = res.breakdown()
    assert res.level == "high" and res.score == 60 and b["excluded_findings"] == 1
    assert [c["rule_id"] for c in b["contributions"]] == ["IN-LOCKIN-24M", "IN-ESCALATION-10", "IN-TERM-ONESIDED"]


# ------------------------------------------------------------------ inputs

def test_docx_word_numbering_headings_and_tables():
    import docx
    d = docx.Document()
    d.add_paragraph("DEALERSHIP AGREEMENT")
    d.add_heading("Payment", level=1)
    d.add_paragraph("The Dealer shall pay each invoice within thirty days of delivery.", style="List Number")
    d.add_paragraph("Late payment shall carry interest at 24% per annum compounded monthly.", style="List Number")
    d.add_heading("Termination", level=1)
    d.add_paragraph("The Company may terminate this agreement on thirty days' notice.")
    t = d.add_table(rows=2, cols=3)
    for (r, c), v in {(0, 0): "Item", (0, 1): "Qty", (0, 2): "Rate", (1, 0): "Pump", (1, 1): "10", (1, 2): "Rs 5,000"}.items():
        t.cell(r, c).text = v
    buf = io.BytesIO()
    d.save(buf)
    prepared = pipeline.prepare_document(buf.getvalue(), "dealer.docx", structured_docx=True)
    text = prepared.parsed.text
    assert "1. The Dealer shall pay" in text and "2. Late payment" in text          # Word auto-numbering restored
    assert "Pump | 10 | Rs 5,000" in text
    heads = [s.heading for s in segment_contract(prepared.parsed, prepared.heading_lines)]
    assert "Payment" in heads and "Termination" in heads
    payment = next(s for s in segment_contract(prepared.parsed, prepared.heading_lines) if s.heading == "Payment")
    assert "24% per annum" in payment.text                                          # list items stay in their clause


def test_prepare_document_still_unpacks_as_a_pair():
    parsed, ocr_used = pipeline.prepare_document(LEASE.encode(), "lease.txt")
    assert parsed.char_count > 100 and ocr_used is False


def _photo(color):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (1200, 1600), color).save(buf, format="JPEG")
    return buf.getvalue()


def test_several_photos_become_pages_in_order_and_are_ocred_page_wise(fake_llm):
    pdf = inputs.photos_to_pdf([_photo("white"), _photo("gray"), _photo("white")])
    parsed = parse_pdf_pages(pdf)
    assert len(parsed.pages) == 3 and parsed.is_scanned

    def ocr(call):
        assert call.files and call.files[0][1] == "application/pdf"
        return json.dumps({"pages": [{"number": i, "text": f"{i}. CLAUSE {i}\nText of page {i} of the contract, read from a photo."}
                                     for i in (1, 2, 3)]})
    state = fake_llm(ocr)
    prepared = pipeline.prepare_document(pdf, "photos.pdf")
    assert prepared.ocr_used and prepared.ocr_calls == 1 and len(state.calls) == 1
    assert [p.number for p in prepared.parsed.pages] == [1, 2, 3]


def test_long_scans_are_ocred_in_page_batches(fake_llm, monkeypatch):
    monkeypatch.setattr(pipeline, "OCR_BATCH_PAGES", 2)
    pdf = inputs.photos_to_pdf([_photo("white")] * 5)
    assert [first for first, _ in pipeline._split_pdf(pdf, 2)] == [1, 3, 5]

    def ocr(call):
        first = int(re.search(r"is page (\d+) of", call.prompt).group(1))
        n = len(parse_pdf_pages(call.files[0][0]).pages)
        return json.dumps({"pages": [{"number": first + i, "text": f"Page {first + i} text of the scanned agreement here."}
                                     for i in range(n)]})
    state = fake_llm(ocr)
    prepared = pipeline.prepare_document(pdf, "scan.pdf")
    assert prepared.ocr_calls == 3 and len(state.calls) == 3
    assert [p.number for p in prepared.parsed.pages] == [1, 2, 3, 4, 5]


def test_hindi_clause_headings_and_devanagari_numbers():
    text = ("अनुबंध\nधारा १ परिभाषाएँ\nइस अनुबंध में क्रेता का अर्थ खरीदने वाला पक्ष है।\n"
            "धारा २ भुगतान\nक्रेता चालान की तारीख से 30 दिनों में भुगतान करेगा।\n"
            "३. समाप्ति\nकोई भी पक्ष 30 दिन की सूचना देकर अनुबंध समाप्त कर सकता है।")
    segs = segment_contract(_doc(text))
    assert [s.number for s in segs if s.number] == ["१", "२", "३"]
    assert any("भुगतान" in s.heading for s in segs)


# ------------------------------------------------------------------ chat helpers

def _row(title, sev, seg, order, quote="q" * 30, ref=None, cat="other", counts=True, status="open", source="rule"):
    return SimpleNamespace(title=title, severity=sev, segment_id=seg, sort_order=order, quote=quote, verified=True, page=None,
                           legal_ref=ref, legal_note="note" if ref else None, category=cat, counts_toward_score=counts,
                           status=status, source=source)


def test_negotiation_order_is_severity_then_document_order():
    rows = [_row("Medium thing", "medium", "S2", 0), _row("Critical thing", "critical", "S5", 1),
            _row("High thing", "high", "S3", 2), _row("Dismissed", "critical", "S9", 3, status="dismissed"),
            _row("Unverified", "critical", "S8", 4, counts=False)]
    ranked = service.rank_findings(rows)
    assert [f.title for f in ranked] == ["Critical thing", "High thing", "Medium thing"]
    cites = service.with_top_findings([], ranked)
    assert [c["finding"] for c in cites] == ["Critical thing", "High thing", "Medium thing"]


def test_enforceability_question_gets_the_matching_legal_note():
    ranked = [_row("Non-compete after exit", "high", "S4", 0, ref="Indian Contract Act 1872 s.27", cat="non_compete"),
              _row("Penalty", "high", "S6", 1, ref="Indian Contract Act 1872 s.74", cat="penalty_damages")]
    notes = service.relevant_notes(ranked, [], "Is the non-compete enforceable?")
    assert [n.legal_ref for n in notes] == ["Indian Contract Act 1872 s.27"]
    assert "s.27" in service.legal_notes_text(notes, hindi=False)
    assert "कानूनी नोट" in service.legal_notes_text(notes, hindi=True)


def test_long_contract_chat_sends_the_relevant_sections():
    segs = segment_contract(_doc(LEASE))
    budget = sum(len(s.text) + len(s.heading) + 40 for s in segs) // 3
    chosen, selected = service.select_for_question(segs, "When is the security deposit refunded?", set(), budget)
    assert selected and any("refund the security deposit" in s.text for s in chosen)
    assert [s.id for s in chosen] == sorted((s.id for s in chosen), key=lambda i: int(i[1:]))   # document order
