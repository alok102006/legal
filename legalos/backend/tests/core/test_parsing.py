import io

import pytest

from app.core.rag import parsing
from app.shared.exceptions import BadRequestException
from app.workspaces.contract_intelligence.engine.segmenter import segment_contract
from tests.workspaces.contract_intelligence.sample_contracts import SUPPLY_AGREEMENT


def _text_pdf(pages):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for page_lines in pages:
        y = 800
        for line in page_lines:
            c.drawString(40, y, line[:110])
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def _image_only_pdf():
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    img = Image.new("RGB", (600, 800), "white")
    ImageDraw.Draw(img).text((40, 40), "scanned page", fill="black")
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawImage(ImageReader(img), 0, 0, width=595, height=842)
    c.showPage()
    c.save()
    return buf.getvalue()


def test_text_pdf_is_parsed_with_page_numbers():
    lines = [l for l in SUPPLY_AGREEMENT.splitlines() if l.strip()]
    half = len(lines) // 2
    pdf = _text_pdf([lines[:half], lines[half:]])
    doc = parsing.parse_contract_file(pdf, "agreement.pdf")
    assert doc.kind == "pdf" and doc.has_pages and not doc.is_scanned
    assert [p.number for p in doc.pages] == [1, 2]
    segs = segment_contract(doc)
    assert any(s.page_start == 2 for s in segs) and any(s.page_start == 1 for s in segs)


def test_scanned_pdf_is_flagged():
    doc = parsing.parse_contract_file(_image_only_pdf(), "scan.pdf")
    assert doc.is_scanned and any("scanned" in w.lower() for w in doc.warnings)


def test_corrupt_pdf_gives_friendly_error():
    with pytest.raises(BadRequestException):
        parsing.parse_contract_file(b"%PDF-1.4 definitely not a pdf", "bad.pdf")


def test_docx_roundtrip_keeps_order_and_tables():
    import docx

    d = docx.Document()
    d.add_paragraph("1. PAYMENT")
    d.add_paragraph("Pay within 30 days.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Item"
    t.rows[0].cells[1].text = "Rs. 5,00,000"
    d.add_paragraph("2. TERM")
    buf = io.BytesIO()
    d.save(buf)
    doc = parsing.parse_contract_file(buf.getvalue(), "a.docx")
    text = doc.text
    assert text.index("PAYMENT") < text.index("Rs. 5,00,000") < text.index("TERM")
    assert not doc.has_pages


def test_text_encoding_fallback():
    doc = parsing.parse_contract_file("Rs. ₹5,00,000 – payable".encode("cp1252", errors="ignore"), "a.txt")
    assert "payable" in doc.text


def test_unsupported_extension():
    with pytest.raises(BadRequestException):
        parsing.parse_contract_file(b"x", "a.exe")


# ------------------------------------------------------------------ general company documents

def _docx_bytes(build) -> bytes:
    import docx

    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_structured_keeps_headings_lists_and_tables_in_order():
    def build(d):
        d.add_heading("Leave Policy", level=1)
        d.add_paragraph("Employees must carry:")
        d.add_paragraph("ID card", style="List Bullet")
        d.add_paragraph("Safety shoes", style="List Bullet")
        d.add_heading("Leave entitlement", level=2)
        t = d.add_table(rows=2, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Type", "Days"
        t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Casual", "12"
        d.add_paragraph("Unused leave lapses on 31 March.")

    doc = parsing.parse_company_document(_docx_bytes(build), "policy.docx")
    text = doc.text
    assert "# Leave Policy" in text and "## Leave entitlement" in text
    assert "• ID card" in text and "• Safety shoes" in text
    assert "Type | Days\nCasual | 12" in text
    assert text.index("# Leave Policy") < text.index("• ID card") < text.index("Type | Days") < text.index("Unused leave")


def _xlsx_bytes() -> bytes:
    import datetime

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["Dealer price list FY 2026-27"])
    ws.append([])
    ws.append(["SKU", "Description", "Price (Rs)", "Valid from"])
    ws.append(["SIP-FAS-003", "Hex nut M10", 310.0, datetime.date(2026, 4, 1)])
    ws.append(["SIP-BRG-102", "Ball bearing", 285.5, None])
    ws2 = wb.create_sheet("Contacts")
    ws2.append(["Name", "Phone"])
    ws2.append(["Hitesh Vyas", "98250 00000"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_rows_become_self_describing_lines_per_sheet():
    doc = parsing.parse_company_document(_xlsx_bytes(), "prices.xlsx")
    text = doc.text
    assert doc.kind == "sheet" and not doc.has_pages
    assert "# Sheet: Prices" in text and "# Sheet: Contacts" in text
    assert "Dealer price list FY 2026-27" in text                                  # title row kept as text
    assert "SKU: SIP-FAS-003 | Description: Hex nut M10 | Price (Rs): 310 | Valid from: 01 Apr 2026" in text
    assert "Price (Rs): 285.5" in text and "Valid from:" not in text.split("SIP-BRG-102")[1].split("\n")[0]
    assert "Name: Hitesh Vyas | Phone: 98250 00000" in text


def test_csv_with_semicolons_and_utf16_is_read():
    data = "Item;Rate\nCotton yarn;Rs 285 per kg\nसूती धागा;₹ 290\n".encode("utf-16")
    doc = parsing.parse_company_document(data, "rates.csv")
    assert "Item: Cotton yarn | Rate: Rs 285 per kg" in doc.text and "Item: सूती धागा | Rate: ₹ 290" in doc.text


@pytest.mark.parametrize("name,data,reason", [
    ("a.docx", b"\xd0\xcf\x11\xe0" + b"\x00" * 100, "encrypted"),
    ("a.xlsx", b"\xd0\xcf\x11\xe0" + b"\x00" * 100, "encrypted"),
    ("a.xlsx", b"PK\x03\x04 not really a workbook", "corrupt"),
    ("a.docx", b"PK\x03\x04 not really a document", "corrupt"),
    ("a.txt", b"   \n\n  ", "empty"),
    ("a.csv", b"", "empty"),
    ("a.exe", b"MZ", "unsupported"),
])
def test_unreadable_files_raise_a_typed_error(name, data, reason):
    with pytest.raises(parsing.DocumentParseError) as e:
        parsing.parse_company_document(data, name)
    assert e.value.reason == reason and isinstance(e.value, BadRequestException) and e.value.detail


def test_too_many_pages_message_says_what_to_do(monkeypatch):
    monkeypatch.setattr(parsing, "MAX_PAGES", 1)
    with pytest.raises(parsing.DocumentParseError) as e:
        parsing.parse_company_document(_text_pdf([["page one"], ["page two"]]), "big.pdf")
    assert e.value.reason == "too_long" and "2 pages" in e.value.detail and "Split it" in e.value.detail


def test_encrypted_pdf_is_reported_as_password_protected():
    pypdf = pytest.importorskip("pypdf")
    writer = pypdf.PdfWriter(clone_from=io.BytesIO(_text_pdf([["secret text"]])))
    writer.encrypt("pw")
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(parsing.DocumentParseError) as e:
        parsing.parse_company_document(buf.getvalue(), "locked.pdf")
    assert e.value.reason == "encrypted"


def _table_pdf() -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    table = Table([["SKU", "Description", "Price"], ["SIP-FAS-003", "Hex nut M10", "310"], ["SIP-BRG-102", "Ball bearing", "285"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    styles = getSampleStyleSheet()
    SimpleDocTemplate(buf, pagesize=A4).build([Paragraph("Price list effective 1 April 2026.", styles["Normal"]), table,
                                                Paragraph("Rates exclude GST.", styles["Normal"])])
    return buf.getvalue()


def test_pdf_ruled_tables_become_pipe_rows_in_reading_order():
    doc = parsing.parse_company_document(_table_pdf(), "prices.pdf")
    text = doc.text
    assert "SKU | Description | Price" in text and "SIP-FAS-003 | Hex nut M10 | 310" in text
    assert text.index("Price list effective") < text.index("SKU | Description") < text.index("Rates exclude GST")
    assert parsing.parse_pdf_pages(_table_pdf()).text.count("|") == 0            # contract parsing is unchanged


def test_image_only_pages_of_a_typed_pdf_are_listed_for_ocr():
    from PIL import Image
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in ["This leave policy applies to all permanent employees of the company.", "Casual leave is twelve days a year.",
                 "Sick leave is seven days a year and needs a doctor's certificate beyond three days.",
                 "Annexure A (the signed approval) follows on the next page."]:
        c.drawString(40, y, line)
        y -= 16
    c.showPage()
    c.drawImage(ImageReader(Image.new("RGB", (300, 400), "white")), 0, 0, width=595, height=842)
    c.showPage()
    c.save()
    doc = parsing.parse_company_document(buf.getvalue(), "mixed.pdf")
    assert not doc.is_scanned and doc.image_pages == [2]
