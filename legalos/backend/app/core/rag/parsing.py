"""Page-aware document parsing (PDF, DOCX, TXT/MD).

PDF uses pdfplumber (MIT) - no AGPL dependency. Scanned PDFs (almost no extractable text)
are flagged via `is_scanned` so callers can route them to Gemini OCR.
"""
import io
import logging
import re
from dataclasses import dataclass, field
from typing import List, Optional

from app.shared.exceptions import BadRequestException

logger = logging.getLogger("legalos.parsing")

MAX_PAGES = 150
# Below this average, a PDF is treated as scanned / image-only.
SCANNED_CHARS_PER_PAGE = 60


class DocumentParseError(BadRequestException):
    """A file that cannot be read, with a machine-readable `reason` so callers and the UI can react
    (encrypted | empty | corrupt | too_long | unsupported). Subclasses BadRequestException, so existing
    handlers keep working and the user sees `detail`."""

    def __init__(self, reason: str, detail: str):
        super().__init__(detail)
        self.reason = reason


@dataclass
class Page:
    number: int
    text: str


@dataclass
class ParsedDocument:
    pages: List[Page]
    kind: str  # "pdf" | "docx" | "text" | "sheet" | "image"
    has_pages: bool  # False for DOCX/TXT where real page numbers are unknown
    is_scanned: bool = False
    warnings: List[str] = field(default_factory=list)
    image_pages: List[int] = field(default_factory=list)  # PDF pages with (almost) no text but an image: need OCR

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Strip control characters / private-use glyphs that PDFs often contain.
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f-�]", "", text)
    return text.strip()


def parse_pdf_pages(file_bytes: bytes) -> ParsedDocument:
    return _parse_pdf(file_bytes, with_tables=False)


def parse_pdf_with_tables(file_bytes: bytes) -> ParsedDocument:
    """Like `parse_pdf_pages`, but ruled tables are rendered as `a | b | c` rows in reading order, so a price list
    or schedule keeps its columns (plain text extraction runs the cells together)."""
    return _parse_pdf(file_bytes, with_tables=True)


def _too_many_pages(n: int) -> DocumentParseError:
    return DocumentParseError(
        "too_long",
        f"This PDF has {n} pages; documents of up to {MAX_PAGES} pages are supported. "
        "Split it into smaller parts (for example by chapter or annexure) and upload each part.")


def _parse_pdf(file_bytes: bytes, with_tables: bool) -> ParsedDocument:
    try:
        import pdfplumber  # lazy: heavy import
    except ImportError as e:  # pragma: no cover
        raise BadRequestException("PDF support is not installed on the server.") from e

    pages: List[Page] = []
    image_pages: List[int] = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            if len(pdf.pages) > MAX_PAGES:
                raise _too_many_pages(len(pdf.pages))
            for i, page in enumerate(pdf.pages, 1):
                raw = _page_text_with_tables(page) if with_tables else (page.extract_text() or "")
                text = _normalize(raw)
                pages.append(Page(i, text))
                if len(text) < SCANNED_CHARS_PER_PAGE and page.images:
                    image_pages.append(i)
    except BadRequestException:
        raise
    except Exception as e:
        msg = f"{type(e).__name__} {e}".lower()
        if "password" in msg or "encrypt" in msg:
            raise DocumentParseError("encrypted", "This PDF is password-protected. Please upload an unlocked copy.")
        raise DocumentParseError("corrupt", f"This PDF is damaged or incomplete and could not be read ({str(e)[:120] or type(e).__name__}). "
                                            "Re-download or re-export it and upload again.")

    if not pages:
        raise DocumentParseError("empty", "The PDF has no pages.")
    total = sum(len(p.text) for p in pages)
    scanned = (total / len(pages)) < SCANNED_CHARS_PER_PAGE
    doc = ParsedDocument(pages=pages, kind="pdf", has_pages=True, is_scanned=scanned, image_pages=image_pages)
    if scanned:
        doc.warnings.append("This looks like a scanned PDF with little selectable text; OCR is required.")
    return doc


def _table_lines(rows: List[List[Optional[str]]]) -> List[str]:
    """pdfplumber table cells -> `a | b | c` lines (merged cells come back as None and are skipped)."""
    out = []
    for row in rows:
        cells = [" ".join((c or "").split()).replace("|", "/") for c in row]
        if any(cells):
            out.append(" | ".join(c for c in cells if c))
    return out


def _page_text_with_tables(page) -> str:
    """Page text in reading order with each ruled table rendered as pipe rows. Falls back to plain extraction on
    any layout problem: a slightly worse rendering beats a failed upload."""
    try:
        tables = sorted((t for t in page.find_tables() if len(t.rows) >= 2), key=lambda t: t.bbox[1])
        if not tables:
            return page.extract_text() or ""
        x0, top0, x1, bottom0 = page.bbox
        parts: List[str] = []
        cursor = top0
        for t in tables:
            tx0, ttop, tx1, tbottom = t.bbox
            if ttop < cursor:                         # overlapping detections: keep the first
                continue
            if ttop > cursor:
                parts.append(page.crop((x0, cursor, x1, ttop)).extract_text() or "")
            lines = _table_lines(t.extract())
            if len(lines) >= 2:
                band = page.crop((x0, ttop, x1, tbottom))
                beside = band.filter(lambda o: not (o.get("object_type") == "char" and tx0 <= o["x0"] and o["x1"] <= tx1)).extract_text() or ""
                parts.append(beside)
                parts.append("\n" + "\n".join(lines) + "\n")
            else:
                parts.append(page.crop((x0, ttop, x1, tbottom)).extract_text() or "")
            cursor = tbottom
        if cursor < bottom0:
            parts.append(page.crop((x0, cursor, x1, bottom0)).extract_text() or "")
        return "\n".join(p for p in parts if p.strip())
    except Exception as e:                             # pragma: no cover - depends on odd PDFs
        logger.info("table-aware extraction failed, using plain text: %s", e)
        return page.extract_text() or ""


def parse_docx_ordered(file_bytes: bytes) -> ParsedDocument:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise BadRequestException(f"Failed to read DOCX: {str(e)[:200]}")

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: List[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, document).text.strip()
            if t:
                lines.append(t)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = []
                for c in row.cells:
                    ct = c.text.strip()
                    if ct and (not cells or cells[-1] != ct):  # merged cells repeat
                        cells.append(ct)
                if cells:
                    lines.append(" | ".join(cells))
    text = _normalize("\n".join(lines))
    return ParsedDocument(pages=[Page(1, text)], kind="docx", has_pages=False)


def parse_text_bytes(file_bytes: bytes) -> ParsedDocument:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = file_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        text = file_bytes.decode("utf-8", errors="ignore")
    return ParsedDocument(pages=[Page(1, _normalize(text))], kind="text", has_pages=False)


# ------------------------------------------------------------------ general company documents (Documents workspace)
#
# The functions below serve "any document an SME uploads" (policies, price lists, registers, certificates). They
# are additive: contract parsing above keeps its exact behaviour. Conventions shared with the chunker:
#   - `# Heading` lines (markdown style) mark real headings found in the file's structure (DOCX styles, sheet names);
#   - tables are `a | b | c` lines, the first line of a block being the header row;
#   - spreadsheet rows are self-describing `Header: value | Header: value` lines, so one row is one retrievable fact.

MAX_SHEET_ROWS = 10_000      # across all sheets; more than this is a database export, not a document
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"


def decode_text(file_bytes: bytes) -> str:
    """Bytes -> text for TXT/MD/CSV: honours UTF-8/UTF-16 BOMs (Excel's "Unicode text" export is UTF-16), then
    UTF-8, then Windows-1252 (common for files saved by older Indian accounting software)."""
    if file_bytes.startswith((b"\xff\xfe", b"\xfe\xff")):
        return file_bytes.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("latin-1")


def _docx_heading_level(style_name: str) -> Optional[int]:
    name = (style_name or "").lower()
    if name == "title":
        return 1
    m = re.match(r"heading\s*(\d)", name)
    return min(int(m.group(1)), 6) if m else None


def parse_docx_structured(file_bytes: bytes) -> ParsedDocument:
    """DOCX -> text in document order: heading-styled paragraphs become `# Heading` lines (level from the style),
    list paragraphs become `• item`, tables become pipe rows (header first) separated by blank lines."""
    if file_bytes.startswith(_OLE_MAGIC):
        raise DocumentParseError("encrypted", "This Word file is password-protected or in the old .doc format. "
                                              "Remove the password or save it as .docx, then upload again.")
    import docx  # python-docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise DocumentParseError("corrupt", f"This Word file could not be opened ({str(e)[:120]}). Re-save it as .docx and upload again.")

    lines: List[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = " ".join(para.text.split())
            if not text:
                continue
            level = _docx_heading_level(para.style.name if para.style is not None else "")
            is_list = child.pPr is not None and child.pPr.numPr is not None or "list" in (para.style.name if para.style is not None else "").lower()
            if level and len(text) <= 150:
                lines += ["", "#" * level + " " + text, ""]
            elif is_list and not re.match(r"^([-•*]|\(?\w{1,3}[.)])\s", text):
                lines.append("• " + text)
            else:
                lines.append(text)
        elif tag == "tbl":
            rows = []
            for row in Table(child, document).rows:
                cells: List[str] = []
                for c in row.cells:
                    ct = " ".join(c.text.split()).replace("|", "/")
                    if ct and (not cells or cells[-1] != ct):     # merged cells repeat their text
                        cells.append(ct)
                if cells:
                    rows.append(" | ".join(cells))
            lines += [""] + rows + [""]
    text = _normalize("\n".join(lines))
    if not text:
        raise DocumentParseError("empty", "This Word file has no text in it.")
    return ParsedDocument(pages=[Page(1, text)], kind="docx", has_pages=False)


def _cell(value) -> str:
    """Spreadsheet cell -> display text the way a person would read it (no '12.0', no '00:00:00')."""
    import datetime as dt

    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dt.datetime):
        return value.strftime("%d %b %Y") if value.time() == dt.time() else value.strftime("%d %b %Y %H:%M")
    if isinstance(value, dt.date):
        return value.strftime("%d %b %Y")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.4f}".rstrip("0").rstrip(".")
    return " ".join(str(value).split()).replace("|", "/")


def _is_header(row: List[str]) -> bool:
    """A header row has at least two filled cells, mostly text, none of them long sentences."""
    filled = [c for c in row if c]
    if len(filled) < 2:
        return False
    texty = [c for c in filled if not re.fullmatch(r"[\d.,%₹$/\- :]+", c)]
    return len(texty) >= 0.6 * len(filled) and all(len(c) <= 60 for c in filled)


def render_rows(rows: List[List[str]]) -> List[str]:
    """Rows of cells -> lines. Lines before the header row (sheet titles, notes) stay plain text; after the header
    each row becomes `Header: value | Header: value` so it is a self-contained, retrievable fact. Without any
    header row the cells are joined with pipes."""
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    keep = [j for j in range(width) if any(j < len(r) and r[j] for r in rows)]          # drop empty columns
    rows = [[r[j] if j < len(r) else "" for j in keep] for r in rows]
    h = next((i for i, r in enumerate(rows[:10]) if _is_header(r)), None)
    if h is None:
        return [" | ".join(c for c in r if c) for r in rows]
    header = [c or f"Column {j + 1}" for j, c in enumerate(rows[h])]
    out = [" ".join(c for c in r if c) for r in rows[:h]]
    for r in rows[h + 1:]:
        pairs = [f"{header[j]}: {c}" for j, c in enumerate(r) if c]
        if pairs:
            out.append(" | ".join(pairs))
    return out


def parse_spreadsheet(file_bytes: bytes, filename: str) -> ParsedDocument:
    """XLSX (every sheet) or CSV -> one text page; each sheet starts with a `# Sheet: name` heading."""
    ext = filename.lower().rsplit(".", 1)[-1]
    sheets: List[tuple[str, List[List[str]]]] = []
    truncated = False
    if ext == "csv":
        import csv

        text = decode_text(file_bytes)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [[_cell(c) for c in r] for r in csv.reader(io.StringIO(text), dialect)]
        truncated = len(rows) > MAX_SHEET_ROWS
        sheets.append((filename.rsplit("/", 1)[-1], rows[:MAX_SHEET_ROWS]))
    else:
        if file_bytes.startswith(_OLE_MAGIC):
            raise DocumentParseError("encrypted", "This Excel file is password-protected or in the old .xls format. "
                                                  "Remove the password or save it as .xlsx, then upload again.")
        import openpyxl

        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as e:
            raise DocumentParseError("corrupt", f"This Excel file could not be opened ({str(e)[:120]}). Re-save it as .xlsx and upload again.")
        budget = MAX_SHEET_ROWS
        for ws in wb.worksheets:
            rows = []
            for values in ws.iter_rows(values_only=True):
                if budget <= 0:
                    truncated = True
                    break
                rows.append([_cell(v) for v in values])
                budget -= 1
            sheets.append((ws.title, rows))
        wb.close()

    parts = []
    for name, rows in sheets:
        lines = render_rows(rows)
        if lines:
            parts.append(f"# Sheet: {name}\n" + "\n".join(lines))
    text = _normalize("\n\n".join(parts))
    if not text:
        raise DocumentParseError("empty", "The spreadsheet has no data in it.")
    doc = ParsedDocument(pages=[Page(1, text)], kind="sheet", has_pages=False)
    if truncated:
        doc.warnings.append(f"Only the first {MAX_SHEET_ROWS:,} rows were read.")
    return doc


def parse_company_document(file_bytes: bytes, filename: str) -> ParsedDocument:
    """Any common office file -> ParsedDocument (images and scanned pages are OCR'd by the caller, which owns the
    AI client). Raises DocumentParseError with a reason the UI can show."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if not file_bytes:
        raise DocumentParseError("empty", "The file is empty.")
    if ext == "pdf":
        return parse_pdf_with_tables(file_bytes)
    if ext == "docx":
        return parse_docx_structured(file_bytes)
    if ext in ("xlsx", "csv"):
        return parse_spreadsheet(file_bytes, filename)
    if ext in ("txt", "md"):
        text = _normalize(decode_text(file_bytes))
        if not text:
            raise DocumentParseError("empty", "The file has no text in it.")
        return ParsedDocument(pages=[Page(1, text)], kind="text", has_pages=False)
    raise DocumentParseError("unsupported", f"'.{ext}' files are not supported. Use PDF, Word, Excel, CSV, text or image files.")


def parse_contract_file(file_bytes: bytes, filename: str) -> ParsedDocument:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        return parse_pdf_pages(file_bytes)
    if ext in ("docx",):
        return parse_docx_ordered(file_bytes)
    if ext in ("txt", "md"):
        return parse_text_bytes(file_bytes)
    raise BadRequestException(f"Unsupported file type '.{ext}'. Supported: PDF, DOCX, TXT, MD.")
